"""Turn raw documents (PDF, DOCX, HTML, MD/TXT) into per-page text.

Scanned PDF pages contain no text layer, so pages with almost no extractable
text are rendered to an image and read with Tesseract OCR.
"""
import html
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

import pypdfium2
import pytesseract
from docx import Document
from pypdf import PdfReader

MIN_CHARS_BEFORE_OCR = 200
OCR_RENDER_SCALE = 2.5

_DEFAULT_TESSERACT_PATH = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
if not shutil.which("tesseract") and _DEFAULT_TESSERACT_PATH.exists():
    pytesseract.pytesseract.tesseract_cmd = str(_DEFAULT_TESSERACT_PATH)


@dataclass
class Page:
    text: str
    page_number: int | None  # None when the source has no page concept (e.g. .md)


def extract_pages(file_path: Path) -> list[Page]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix in (".html", ".htm"):
        return _extract_html(file_path)
    if suffix in (".md", ".txt"):
        return _extract_plain_text(file_path)
    raise ValueError(f"Unsupported file type: {file_path.name}")


def _extract_pdf(file_path: Path) -> list[Page]:
    reader = PdfReader(str(file_path))
    rendered = None
    pages = []

    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if len(text.strip()) < MIN_CHARS_BEFORE_OCR:
            rendered = rendered or pypdfium2.PdfDocument(str(file_path))
            print(f"  OCR {file_path.name} page {i + 1}/{len(reader.pages)}", flush=True)
            ocr_text = _ocr_page(rendered[i])
            text = ocr_text if len(ocr_text) > len(text) else text
        if text.strip():
            pages.append(Page(text=text, page_number=i + 1))
    return pages


def _ocr_page(pdfium_page) -> str:
    image = pdfium_page.render(scale=OCR_RENDER_SCALE).to_pil()
    return pytesseract.image_to_string(image)


def _extract_docx(file_path: Path) -> list[Page]:
    doc = Document(str(file_path))
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Page(text=text, page_number=None)] if text.strip() else []


def _extract_html(file_path: Path) -> list[Page]:
    raw = file_path.read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", "", raw, flags=re.DOTALL | re.IGNORECASE)
    raw = re.sub(r"</(p|div|tr|li|h\d)>|<br\s*/?>", "\n", raw, flags=re.IGNORECASE)
    text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
    return [Page(text=text, page_number=None)] if text.strip() else []


def _extract_plain_text(file_path: Path) -> list[Page]:
    text = file_path.read_text(encoding="utf-8")
    return [Page(text=text, page_number=None)] if text.strip() else []
